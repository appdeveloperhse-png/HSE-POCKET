__version__ = "1.0.0"

import os
import sqlite3
from datetime import datetime

from kivy.app import App
from kivy.core.window import Window
from kivy.metrics import dp
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.button import Button
from kivy.uix.gridlayout import GridLayout
from kivy.uix.label import Label
from kivy.uix.popup import Popup
from kivy.uix.scrollview import ScrollView
from kivy.uix.spinner import Spinner
from kivy.uix.tabbedpanel import TabbedPanel
from kivy.uix.tabbedpanel import TabbedPanelHeader
from kivy.uix.textinput import TextInput

import openpyxl
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


# ============================================================
# APPLICATION PATHS
# ============================================================

APP_DIR = os.path.join(
    App.get_running_app().user_data_dir
    if App.get_running_app()
    else os.path.expanduser("~"),
    "HSE_Management"
)

os.makedirs(APP_DIR, exist_ok=True)

DB_NAME = os.path.join(APP_DIR, "hse_management.db")
EXPORT_DIR = os.path.join(APP_DIR, "exports")

os.makedirs(EXPORT_DIR, exist_ok=True)


# ============================================================
# DATABASE
# ============================================================

def get_connection():
    return sqlite3.connect(DB_NAME)


def init_db():

    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS observations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            doc_no TEXT,
            category TEXT,
            obs_type TEXT,
            priority TEXT,
            description TEXT,
            action TEXT,
            target_date TEXT,
            status TEXT,
            created_at TEXT
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS incidents (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            doc_no TEXT,
            title TEXT,
            date_time TEXT,
            location TEXT,
            rca_method TEXT,
            rca_details TEXT
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS audits (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            doc_no TEXT,
            audit_type TEXT,
            iso_standard TEXT,
            finding_type TEXT,
            details TEXT,
            created_at TEXT
        )
    """)

    cursor.execute("""
        CREATE TABLE IF NOT EXISTS capa (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            source TEXT,
            action_item TEXT,
            assignee TEXT,
            target_date TEXT,
            status TEXT,
            created_at TEXT
        )
    """)

    conn.commit()
    conn.close()


# ============================================================
# GENERAL HELPERS
# ============================================================

def timestamp():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def document_number(prefix):
    return f"{prefix}-{datetime.now().strftime('%Y%m%d%H%M%S')}"


def safe_filename(value):
    characters = '<>:"/\\|?*'
    for char in characters:
        value = value.replace(char, "_")
    return value


# ============================================================
# WORD EXPORT
# ============================================================

def export_stop_card_docx(
    doc_no,
    category,
    obs_type,
    priority,
    description,
    action
):

    filename = safe_filename(
        f"STOP_Card_{doc_no}.docx"
    )

    path = os.path.join(EXPORT_DIR, filename)

    doc = Document()

    doc.add_heading(
        "HSE OBSERVATION & STOP CARD",
        level=1
    )

    doc.add_paragraph(
        f"Document No: {doc_no}"
    )

    doc.add_paragraph(
        f"Category: {category}"
    )

    doc.add_paragraph(
        f"Observation Type: {obs_type}"
    )

    doc.add_paragraph(
        f"Priority: {priority}"
    )

    doc.add_heading(
        "Observation",
        level=2
    )

    doc.add_paragraph(description)

    doc.add_heading(
        "Corrective Action",
        level=2
    )

    doc.add_paragraph(action)

    doc.add_paragraph(
        f"Generated: {timestamp()}"
    )

    doc.save(path)

    return path


# ============================================================
# PDF EXPORT
# ============================================================

def export_stop_card_pdf(
    doc_no,
    category,
    obs_type,
    priority,
    description,
    action
):

    filename = safe_filename(
        f"STOP_Card_{doc_no}.pdf"
    )

    path = os.path.join(EXPORT_DIR, filename)

    pdf = canvas.Canvas(path, pagesize=A4)

    width, height = A4

    y = height - 50

    pdf.setFont("Helvetica-Bold", 16)

    pdf.drawString(
        40,
        y,
        "HSE OBSERVATION & STOP CARD"
    )

    y -= 40

    pdf.setFont("Helvetica", 10)

    data = [
        f"Document No: {doc_no}",
        f"Category: {category}",
        f"Observation Type: {obs_type}",
        f"Priority: {priority}",
        "",
        "Observation:",
        description,
        "",
        "Corrective Action:",
        action,
        "",
        f"Generated: {timestamp()}"
    ]

    for line in data:

        if y < 50:
            pdf.showPage()
            y = height - 50

        if line == "":
            y -= 12
            continue

        lines = line.split("\n")

        for text_line in lines:

            if y < 50:
                pdf.showPage()
                y = height - 50

            pdf.drawString(
                40,
                y,
                text_line[:110]
            )

            y -= 16

    pdf.save()

    return path


# ============================================================
# EXCEL EXPORT
# ============================================================

def export_capa_excel():

    filename = "CAPA_Register.xlsx"

    path = os.path.join(
        EXPORT_DIR,
        filename
    )

    conn = get_connection()

    cursor = conn.cursor()

    cursor.execute("""
        SELECT
            id,
            source,
            action_item,
            assignee,
            target_date,
            status,
            created_at
        FROM capa
        ORDER BY id DESC
    """)

    rows = cursor.fetchall()

    conn.close()

    workbook = openpyxl.Workbook()

    sheet = workbook.active

    sheet.title = "CAPA Register"

    sheet.append([
        "ID",
        "Source",
        "Action Item",
        "Assignee",
        "Target Date",
        "Status",
        "Created At"
    ])

    for row in rows:
        sheet.append(list(row))

    sheet.freeze_panes = "A2"

    workbook.save(path)

    return path


# ============================================================
# MAIN APPLICATION
# ============================================================

class HSEApp(App):

    def build(self):

        self.title = "HSE Management System"

        Window.softinput_mode = "below_target"

        init_db()

        main_layout = BoxLayout(
            orientation="vertical"
        )

        header = Label(
            text="HSE MANAGEMENT SYSTEM",
            size_hint_y=None,
            height=dp(55),
            font_size="18sp",
            bold=True
        )

        main_layout.add_widget(header)

        self.tabs = TabbedPanel(
            do_default_tab=False,
            tab_width=dp(130)
        )

        # Dashboard

        dashboard_tab = TabbedPanelHeader(
            text="Dashboard"
        )

        dashboard_tab.content = (
            self.create_dashboard()
        )

        self.tabs.add_widget(
            dashboard_tab
        )

        # Observations

        observation_tab = TabbedPanelHeader(
            text="Observations"
        )

        observation_tab.content = (
            self.create_observation_view()
        )

        self.tabs.add_widget(
            observation_tab
        )

        # Incidents

        incident_tab = TabbedPanelHeader(
            text="Incidents"
        )

        incident_tab.content = (
            self.create_incident_view()
        )

        self.tabs.add_widget(
            incident_tab
        )

        # Audits

        audit_tab = TabbedPanelHeader(
            text="Audits"
        )

        audit_tab.content = (
            self.create_audit_view()
        )

        self.tabs.add_widget(
            audit_tab
        )

        # CAPA

        capa_tab = TabbedPanelHeader(
            text="CAPA"
        )

        capa_tab.content = (
            self.create_capa_view()
        )

        self.tabs.add_widget(
            capa_tab
        )

        main_layout.add_widget(
            self.tabs
        )

        return main_layout


# ============================================================
# DASHBOARD
# ============================================================

    def create_dashboard(self):

        layout = BoxLayout(
            orientation="vertical",
            padding=dp(15),
            spacing=dp(12)
        )

        self.dashboard_label = Label(
            text="Loading...",
            font_size="17sp"
        )

        layout.add_widget(
            self.dashboard_label
        )

        refresh = Button(
            text="Refresh Dashboard",
            size_hint_y=None,
            height=dp(50)
        )

        refresh.bind(
            on_release=lambda x:
            self.refresh_dashboard()
        )

        layout.add_widget(refresh)

        self.refresh_dashboard()

        return layout


    def refresh_dashboard(self):

        conn = get_connection()

        cursor = conn.cursor()

        cursor.execute(
            "SELECT COUNT(*) FROM observations"
        )

        observations = cursor.fetchone()[0]

        cursor.execute(
            "SELECT COUNT(*) FROM incidents"
        )

        incidents = cursor.fetchone()[0]

        cursor.execute(
            "SELECT COUNT(*) FROM audits"
        )

        audits = cursor.fetchone()[0]

        cursor.execute("""
            SELECT COUNT(*)
            FROM capa
            WHERE status != 'Closed'
        """)

        open_capa = cursor.fetchone()[0]

        conn.close()

        self.dashboard_label.text = (
            f"Observations / STOP Cards: {observations}\n\n"
            f"Incidents: {incidents}\n\n"
            f"Audit Findings: {audits}\n\n"
            f"Open CAPA: {open_capa}"
        )


# ============================================================
# OBSERVATION
# ============================================================

    def create_observation_view(self):

        scroll = ScrollView()

        grid = GridLayout(
            cols=2,
            spacing=dp(8),
            padding=dp(10),
            size_hint_y=None
        )

        grid.bind(
            minimum_height=grid.setter("height")
        )

        def label(text):
            return Label(
                text=text,
                size_hint_y=None,
                height=dp(45)
            )

        grid.add_widget(
            label("Category")
        )

        self.obs_cat = Spinner(
            text="PPE",
            values=[
                "PPE",
                "Work at Height",
                "Excavation",
                "Lifting",
                "Electrical",
                "Confined Space",
                "Hot Work",
                "Vehicle Safety",
                "Scaffolding",
                "Housekeeping",
                "Environmental",
                "Other"
            ],
            size_hint_y=None,
            height=dp(45)
        )

        grid.add_widget(
            self.obs_cat
        )

        grid.add_widget(
            label("Observation Type")
        )

        self.obs_type = Spinner(
            text="Unsafe Condition",
            values=[
                "Unsafe Act",
                "Unsafe Condition",
                "Good Practice",
                "Near Miss",
                "Environmental"
            ],
            size_hint_y=None,
            height=dp(45)
        )

        grid.add_widget(
            self.obs_type
        )

        grid.add_widget(
            label("Priority")
        )

        self.obs_priority = Spinner(
            text="Medium",
            values=[
                "Low",
                "Medium",
                "High",
                "Critical"
            ],
            size_hint_y=None,
            height=dp(45)
        )

        grid.add_widget(
            self.obs_priority
        )

        grid.add_widget(
            label("Observation")
        )

        self.obs_desc = TextInput(
            multiline=True,
            size_hint_y=None,
            height=dp(120)
        )

        grid.add_widget(
            self.obs_desc
        )

        grid.add_widget(
            label("Corrective Action")
        )

        self.obs_action = TextInput(
            multiline=True,
            size_hint_y=None,
            height=dp(120)
        )

        grid.add_widget(
            self.obs_action
        )

        save = Button(
            text="SAVE + GENERATE REPORTS",
            size_hint_y=None,
            height=dp(55)
        )

        save.bind(
            on_release=self.save_observation
        )

        grid.add_widget(
            save
        )

        grid.add_widget(
            Label(text="")
        )

        scroll.add_widget(grid)

        return scroll


    def save_observation(self, instance):

        if not self.obs_desc.text.strip():

            self.show_popup(
                "Required",
                "Please enter the observation."
            )

            return

        doc_no = document_number(
            "HSE-OBS"
        )

        conn = get_connection()

        cursor = conn.cursor()

        cursor.execute("""
            INSERT INTO observations (
                doc_no,
                category,
                obs_type,
                priority,
                description,
                action,
                target_date,
                status,
                created_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            doc_no,
            self.obs_cat.text,
            self.obs_type.text,
            self.obs_priority.text,
            self.obs_desc.text,
            self.obs_action.text,
            "30 Days",
            "Open",
            timestamp()
        ))

        cursor.execute("""
            INSERT INTO capa (
                source,
                action_item,
                assignee,
                target_date,
                status,
                created_at
            )
            VALUES (?, ?, ?, ?, ?, ?)
        """, (
            doc_no,
            self.obs_action.text,
            "HSE Supervisor",
            "30 Days",
            "Open",
            timestamp()
        ))

        conn.commit()

        conn.close()

        docx = export_stop_card_docx(
            doc_no,
            self.obs_cat.text,
            self.obs_type.text,
            self.obs_priority.text,
            self.obs_desc.text,
            self.obs_action.text
        )

        pdf = export_stop_card_pdf(
            doc_no,
            self.obs_cat.text,
            self.obs_type.text,
            self.obs_priority.text,
            self.obs_desc.text,
            self.obs_action.text
        )

        self.refresh_dashboard()

        self.show_popup(
            "Saved",
            "Observation saved successfully.\n\n"
            f"Word:\n{docx}\n\n"
            f"PDF:\n{pdf}"
        )


# ============================================================
# INCIDENT
# ============================================================

    def create_incident_view(self):

        scroll = ScrollView()

        grid = GridLayout(
            cols=2,
            spacing=dp(8),
            padding=dp(10),
            size_hint_y=None
        )

        grid.bind(
            minimum_height=grid.setter("height")
        )

        grid.add_widget(
            Label(
                text="Incident Title",
                size_hint_y=None,
                height=dp(45)
            )
        )

        self.inc_title = TextInput(
            size_hint_y=None,
            height=dp(45)
        )

        grid.add_widget(
            self.inc_title
        )

        grid.add_widget(
            Label(
                text="Location",
                size_hint_y=None,
                height=dp(45)
            )
        )

        self.inc_loc = TextInput(
            size_hint_y=None,
            height=dp(45)
        )

        grid.add_widget(
            self.inc_loc
        )

        grid.add_widget(
            Label(
                text="RCA Method",
                size_hint_y=None,
                height=dp(45)
            )
        )

        self.inc_rca_method = Spinner(
            text="5-Why Analysis",
            values=[
                "5-Why Analysis",
                "Fishbone / Ishikawa",
                "ICAM Method"
            ],
            size_hint_y=None,
            height=dp(45)
        )

        grid.add_widget(
            self.inc_rca_method
        )

        grid.add_widget(
            Label(
                text="RCA Details",
                size_hint_y=None,
                height=dp(45)
            )
        )

        self.inc_rca_details = TextInput(
            multiline=True,
            size_hint_y=None,
            height=dp(160)
        )

        grid.add_widget(
            self.inc_rca_details
        )

        save = Button(
            text="SAVE INCIDENT",
            size_hint_y=None,
            height=dp(55)
        )

        save.bind(
            on_release=self.save_incident
        )

        grid.add_widget(save)

        scroll.add_widget(grid)

        return scroll


    def save_incident(self, instance):

        if not self.inc_title.text.strip():

            self.show_popup(
                "Required",
                "Enter incident title."
            )

            return

        doc_no = document_number(
            "HSE-INC"
        )

        conn = get_connection()

        cursor = conn.cursor()

        cursor.execute("""
            INSERT INTO incidents (
                doc_no,
                title,
                date_time,
                location,
                rca_method,
                rca_details
            )
            VALUES (?, ?, ?, ?, ?, ?)
        """, (
            doc_no,
            self.inc_title.text,
            timestamp(),
            self.inc_loc.text,
            self.inc_rca_method.text,
            self.inc_rca_details.text
        ))

        conn.commit()

        conn.close()

        self.refresh_dashboard()

        self.show_popup(
            "Saved",
            f"Incident recorded.\n\n{doc_no}"
        )


# ============================================================
# AUDIT
# ============================================================

    def create_audit_view(self):

        scroll = ScrollView()

        grid = GridLayout(
            cols=2,
            spacing=dp(8),
            padding=dp(10),
            size_hint_y=None
        )

        grid.bind(
            minimum_height=grid.setter("height")
        )

        grid.add_widget(
            Label(
                text="Audit Type",
                size_hint_y=None,
                height=dp(45)
            )
        )

        self.audit_type = Spinner(
            text="Internal Audit",
            values=[
                "Internal Audit",
                "External Audit",
                "Certification",
                "Surveillance"
            ],
            size_hint_y=None,
            height=dp(45)
        )

        grid.add_widget(
            self.audit_type
        )

        grid.add_widget(
            Label(
                text="ISO Standard",
                size_hint_y=None,
                height=dp(45)
            )
        )

        self.audit_iso = Spinner(
            text="ISO 45001",
            values=[
                "ISO 45001",
                "ISO 14001"
            ],
            size_hint_y=None,
            height=dp(45)
        )

        grid.add_widget(
            self.audit_iso
        )

        grid.add_widget(
            Label(
                text="Finding",
                size_hint_y=None,
                height=dp(45)
            )
        )

        self.audit_finding = Spinner(
            text="Minor Nonconformity",
            values=[
                "Major Nonconformity",
                "Minor Nonconformity",
                "OFI",
                "Conformity"
            ],
            size_hint_y=None,
            height=dp(45)
        )

        grid.add_widget(
            self.audit_finding
        )

        grid.add_widget(
            Label(
                text="Finding Details",
                size_hint_y=None,
                height=dp(45)
            )
        )

        self.audit_details = TextInput(
            multiline=True,
            size_hint_y=None,
            height=dp(150)
        )

        grid.add_widget(
            self.audit_details
        )

        save = Button(
            text="SAVE AUDIT FINDING",
            size_hint_y=None,
            height=dp(55)
        )

        save.bind(
            on_release=self.save_audit
        )

        grid.add_widget(save)

        scroll.add_widget(grid)

        return scroll


    def save_audit(self, instance):

        doc_no = document_number(
            "HSE-AUD"
        )

        conn = get_connection()

        cursor = conn.cursor()

        cursor.execute("""
            INSERT INTO audits (
                doc_no,
                audit_type,
                iso_standard,
                finding_type,
                details,
                created_at
            )
            VALUES (?, ?, ?, ?, ?, ?)
        """, (
            doc_no,
            self.audit_type.text,
            self.audit_iso.text,
            self.audit_finding.text,
            self.audit_details.text,
            timestamp()
        ))

        conn.commit()

        conn.close()

        self.refresh_dashboard()

        self.show_popup(
            "Saved",
            f"Audit finding recorded.\n\n{doc_no}"
        )


# ============================================================
# CAPA
# ============================================================

    def create_capa_view(self):

        layout = BoxLayout(
            orientation="vertical",
            padding=dp(15),
            spacing=dp(12)
        )

        info = Label(
            text=(
                "CAPA Register\n\n"
                "Observation corrective actions are "
                "automatically added here."
            ),
            font_size="16sp"
        )

        layout.add_widget(info)

        export = Button(
            text="EXPORT CAPA TO EXCEL",
            size_hint_y=None,
            height=dp(55)
        )

        export.bind(
            on_release=self.export_excel
        )

        layout.add_widget(export)

        return layout


    def export_excel(self, instance):

        try:

            path = export_capa_excel()

            self.show_popup(
                "Export Complete",
                f"Excel file created:\n\n{path}"
            )

        except Exception as error:

            self.show_popup(
                "Export Error",
                str(error)
            )


# ============================================================
# POPUP
# ============================================================

    def show_popup(self, title, message):

        content = BoxLayout(
            orientation="vertical",
            padding=dp(10),
            spacing=dp(10)
        )

        text = Label(
            text=message
        )

        close = Button(
            text="OK",
            size_hint_y=None,
            height=dp(50)
        )

        content.add_widget(text)
        content.add_widget(close)

        popup = Popup(
            title=title,
            content=content,
            size_hint=(0.9, 0.55)
        )

        close.bind(
            on_release=popup.dismiss
        )

        popup.open()


# ============================================================
# START
# ============================================================

if __name__ == "__main__":

    HSEApp().run()
